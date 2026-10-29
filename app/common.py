"""Shared helpers: paths, columns, settings, the Tracker (load/save), local AI, pay parsing."""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")  # quiets urllib3/LibreSSL notices on macOS system Python

import openpyxl  # noqa: E402
import requests  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

APP_VERSION = "21"
APP_DIR = Path(__file__).resolve().parent
ROOT = APP_DIR.parent
WB_PATH = ROOT / "Job_Tracker.xlsx"
STATE_PATH = APP_DIR / "state.json"
ACCOUNTS_PATH = APP_DIR / "accounts.json"
BACKUP_DIR = APP_DIR / "backups"
POSTINGS_DIR = ROOT / "Postings"
BACKUPS_KEEP = 15

# ---------------------------------------------------------------- columns
# (key, header, width, kind)  kind: in = value, calc = Excel formula, grp = value in the collapsed group
COLUMNS = [
    ("id", "ID", 5, "in"),
    ("priority", "Priority", 8, "calc"),
    ("company", "Company", 22, "in"),
    ("title", "Title", 30, "in"),
    ("posted", "Posted", 9, "in"),
    ("days_posted", "Days Since Posted", 8, "calc"),
    ("deadline", "Deadline", 9, "in"),
    ("days_left", "Days Left", 7, "calc"),
    ("location", "Location", 20, "in"),
    ("area", "Area", 22, "in"),
    ("work_mode", "Work Mode", 10, "in"),
    ("office_days", "Office Days/Wk", 8, "in"),
    ("pay_min", "Pay Min", 10, "in"),
    ("pay_max", "Pay Max", 10, "in"),
    ("pay_unit", "Pay Unit", 7, "in"),
    ("total_comp", "Est. Total Comp", 12, "calc"),
    ("perks", "Perks", 50, "calc"),
    ("status", "Status", 11, "in"),
    ("applied", "Applied?", 8, "in"),
    ("date_applied", "Date Applied", 10, "in"),
    ("resume", "Resume Version", 24, "in"),
    ("next_action", "Next Action", 28, "in"),
    ("last_update", "Last Update", 10, "in"),
    ("link", "Link", 28, "in"),
    ("contact", "Contact", 20, "in"),
    ("notes", "Notes", 28, "in"),
    ("fit_notes", "Fit Notes", 30, "in"),
    ("source", "Source", 12, "in"),
    ("date_added", "Date Added", 10, "in"),
    ("emp_type", "Employment Type", 12, "in"),
    # ---- collapsed "scoring details" group (click + above the columns to expand)
    ("fit_score", "Fit Score", 8, "grp"),
    ("loc_score", "Location Score", 9, "calc"),
    ("pay_score", "Pay Score", 8, "calc"),
    ("perks_score", "Perks Score", 8, "calc"),
    ("area_type", "Area Type", 9, "grp"),
    ("miles", "Miles", 7, "grp"),
    ("area_bonus", "Area Points", 8, "grp"),
    ("signon", "Sign-on Bonus", 10, "grp"),
    ("bonus_pct", "Bonus Target %", 9, "grp"),
    ("bonus_amt", "Bonus $", 9, "grp"),
    ("k401", "401k Match", 9, "grp"),
    ("relocation", "Relocation", 9, "grp"),
    ("equity", "Equity", 8, "grp"),
    ("esop", "Employee Owned", 9, "grp"),
    ("health", "Strong Health", 9, "grp"),
    ("pto", "PTO Days", 9, "grp"),
    ("gym", "Gym/Wellness", 9, "grp"),
    ("stipend", "WFH Stipend", 9, "grp"),
    ("other_perks", "Other Perks", 24, "grp"),
    ("evidence", "Perk Evidence", 40, "grp"),
    ("description", "Description", 40, "grp"),
    ("email_log", "Email History", 40, "grp"),
    ("geo_key", "Area Checked For", 24, "grp"),
]
KEYS = [c[0] for c in COLUMNS]
COL = {k: i + 1 for i, k in enumerate(KEYS)}
HEADER = {c[0]: c[1] for c in COLUMNS}
CALC = {c[0] for c in COLUMNS if c[3] == "calc"}
GROUP_KEYS = KEYS[KEYS.index("fit_score"):]  # collapsed by default
DATE_KEYS = {"posted", "deadline", "date_applied", "last_update", "date_added"}
HEADER_TO_KEY = {h.lower(): k for k, h, _, _ in COLUMNS}
HEADER_TO_KEY.update({"posted date": "posted", "contact / recruiter": "contact", "email history": "email_log",
                      "priority (0-100)": "priority", "fit score (0-10)": "fit_score"})


def L(key: str) -> str:
    return get_column_letter(COL[key])


JOB_TABS = ["To Apply", "Applied", "Closed"]
TABLE_NAMES = {"To Apply": "tblToApply", "Applied": "tblApplied", "Closed": "tblClosed"}
STATUSES = ["Saved", "Applied", "Screening", "Assessment", "Interview", "Offer",
            "Rejected", "Withdrawn", "Ghosted", "Skipped", "Declined"]
CLOSED = {"Rejected", "Withdrawn", "Ghosted", "Skipped", "Declined"}
STATUS_RANK = {"": 0, "Saved": 0, "Skipped": 0.5, "Applied": 1, "Ghosted": 1, "Screening": 2,
               "Assessment": 3, "Interview": 4, "Rejected": 5, "Withdrawn": 5, "Offer": 6, "Declined": 7}


def status_rank(s) -> float:
    return STATUS_RANK.get(str(s or "").strip().title(), 0)


# ---------------------------------------------------------------- settings
# (key, label, default, help, user_fills, excel_name)
H = "_h"
SETTINGS = [
    (H, "PRIORITY: LOCATION & WORK SETUP  (your preferred areas are on the Areas tab)", None, None, False, None),
    ("home_city", "Home (city, state)", "", "Where you live now. Jobs within the max commute count as Home.", True, None),
    ("home_radius", "Max commute from home (miles)", 40, "Farther than this = you'd have to move.", True, "HOME_RADIUS"),
    ("commute_free", "Full score for commutes up to (miles)", 20, "Past this, Home jobs lose points gradually out to the max commute.", False, "COMMUTE_FREE"),
    ("commute_pen", "Long-commute penalty (points, at max commute, 5 office days)", 2, "Scaled down for fewer office days.", False, "COMMUTE_PEN"),
    ("region_states", "Preferred region (state codes)", "",
     "Places in these states that aren't inside a listed area still rank above 'anywhere else'.", True, None),
    ("region_pen", "Preferred region: points below a listed area", 2, "", False, "REGION_PEN"),
    ("loc_remote", "Score: fully remote", 10, "Each score is 0-10. Office-day scores apply to Home, Areas and Region.", False, "LOC_REMOTE"),
    ("loc_d1", "Score: hybrid, 1 office day/week", 9, "", False, "LOC_D1"),
    ("loc_d2", "Score: hybrid, 2 office days/week", 8, "", False, "LOC_D2"),
    ("loc_d3", "Score: hybrid, 3 office days/week", 6, "", False, "LOC_D3"),
    ("loc_d4", "Score: hybrid, 4 office days/week", 4, "", False, "LOC_D4"),
    ("loc_hybrid", "Score: hybrid, office days not stated", 7, "", False, "LOC_HYBRID"),
    ("loc_onsite", "Score: fully on-site", 3, "", False, "LOC_ONSITE"),
    ("loc_else", "Score: hybrid/on-site anywhere else", 2, "A move you don't really want.", False, "LOC_ELSE"),
    ("relo_area", "Relocation-help bonus: listed areas and region", 2, "Added when the posting offers relocation help.", False, "RELO_AREA"),
    ("relo_else", "Relocation-help bonus: anywhere else", 1, "", False, "RELO_ELSE"),
    ("loc_unknown", "Score: location not listed or not recognized", 5, "Shows '?' in the Area column; fix the Location text.", False, "LOC_UNKNOWN"),
    (H, "PRIORITY: PAY", None, None, False, None),
    ("target_pay", "Target total pay (yearly $)", None, "Est. Total Comp at or above this scores 10/10. Total comp = base + bonus + "
     "401k match + sign-on spread over the years below. Hourly pay x 2,080.", True, "TARGET_PAY"),
    ("min_pay", "Minimum acceptable total pay (yearly $)", None, "At or below this scores 0/10.", True, "MIN_PAY"),
    ("pay_none", "Pay score when pay isn't listed", 5, "0-10", False, "PAY_NONE"),
    ("signon_years", "Spread sign-on bonus over (years)", 2, "", False, "SIGNON_YEARS"),
    (H, "PRIORITY: PERKS (points add up; the Perks score is capped at 10)", None, None, False, None),
    ("p_equity", "Points: equity / stock (RSUs, options, ESPP)", 2, "", False, "P_EQUITY"),
    ("p_esop", "Points: employee-owned / ESOP", 2, "", False, "P_ESOP"),
    ("p_health", "Points: strong health coverage", 2, "Employer pays all or nearly all premiums.", False, "P_HEALTH"),
    ("p_401k", "Points: 401(k) match (any amount)", 1, "A known match % also counts toward total pay.", False, "P_401K"),
    ("p_pto", "Points: generous PTO", 1, "", False, "P_PTO"),
    ("pto_min", "PTO days that count as generous", 20, "Unlimited PTO always counts.", False, "PTO_MIN"),
    ("p_gym", "Points: gym / wellness", 1, "", False, "P_GYM"),
    ("p_stipend", "Points: remote / home-office stipend", 1, "", False, "P_STIPEND"),
    ("p_other", "Points: other notable perk", 1, "e.g. 4-day week, pension, profit sharing, student loan help.", False, "P_OTHER"),
    (H, "PRIORITY: FIT", None, None, False, None),
    ("fit_keywords", "Fit keywords", "", "Specific skills/tools/domains, comma-separated; phrases work (e.g. mechanical design, "
     "SolidWorks, medical device). Title matches count most; benefits/legal text is ignored. 6-15 specific terms beat generic words.", True, None),
    ("my_years", "My years of experience", 2, "Count internships/co-ops. Used to check each posting's 'X+ years required'.", True, None),
    ("dealbreakers", "Dealbreaker words in the job TITLE", "", "Checked against the title only, e.g. Senior, Sr, Staff, Principal, Director, Manager", True, None),
    ("deal_phrases", "Dealbreaker phrases anywhere in the posting", "", "e.g. active security clearance, 100% travel. Leave blank if none.", True, None),
    ("background", "My background (for AI fit)", "", "2-4 sentences: degree, years of experience, strongest skills, roles you want. "
     "Used only by the local AI.", True, None),
    (H, "PRIORITY WEIGHTS (how much each score counts)", None, None, False, None),
    ("w_loc", "Weight: location", 1, "Changing these re-scores every job instantly.", False, "W_LOC"),
    ("w_pay", "Weight: pay", 1, "", False, "W_PAY"),
    ("w_fit", "Weight: fit", 1, "", False, "W_FIT"),
    ("w_perks", "Weight: perks", 0.5, "", False, "W_PERKS"),
    (H, "JOB WATCHER (Check Jobs; companies are on the Watchlist tab)", None, None, False, None),
    ("watch_terms", "Watch: title keywords", "mechanical, civil, structural, electrical, design engineer, test engineer, "
     "manufacturing engineer, process engineer, quality engineer, systems engineer, project engineer, R&D engineer",
     "A new posting needs ALL the words of at least one entry in its title, in any order ('test engineer' also matches "
     "'Engineer, Test'). Put an entry in quotes to require those words together, in order: \"product engineer\" "
     "won't match 'Product Deployment Engineer'. Single words cast a wider net. First 12 are also used as searches "
     "on Workday, Eightfold, Apple and Google.", True, None),
    ("watch_exclude", "Watch: skip titles containing", "intern, internship, co-op, technician, sales, software, "
     "manager, director", "Your title dealbreakers are skipped too.", True, None),
    ("watch_skip_senior", "Watch: skip senior-level titles", "Yes", "Yes/No. On = titles with any word from the two lists "
     "below are skipped by Check Jobs, and score low on fit.", False, None),
    ("watch_senior_words", "Watch: senior-level title words", "Senior, Sr, Snr, Staff, Principal, Prin, Princ, Lead, Leader, "
     "Manager, Mgr, Director, Dir, Head of, Chief, Vice President, VP, SVP, EVP, Distinguished, Fellow",
     "Edit freely. Plurals are matched too ('Lead' also matches 'Leads', not 'Leader').", True, None),
    ("watch_level_words", "Watch: level words (skipped)", "III, IV, V, VI, 3, 4, 5, 6, 7, 8, 9, Level 3, Level 4, Level 5, "
     "Grade 5, P4, P5, P6, L4, L5, L6, E4, E5, E6",
     "Matched as separate words ('Engineer III', 'Engineer 3', 'Level 4'); never inside other words ('V&V', '3D'). A range "
     "that starts below these levels stays in ('Engineer II/III' is kept, 'Engineer III/IV' is skipped).", True, None),
    ("watch_areas", "Watch: locations to include", "Home, Area, Region, Remote, Elsewhere", "Any of: Home, Area, Region, Remote, "
     "Elsewhere (rest of the US), Outside US. Every match is labeled in the Area column. Postings with unclear "
     "locations are always shown.", False, None),
    ("watch_days", "Watch: scan window (days)", 14,
     "How far back a big company's first scan looks, and the most any later scan looks back. After the first scan, "
     "each scan covers the days since that company was last scanned (+2), at least 7. Postings without a date are "
     "always checked.", False, None),
    ("watch_small", "Watch: small company (up to this many postings)", 60,
     "A small company's FIRST scan shows every matching job that's still open, however old (big companies keep "
     "old postings open, so their first scan uses the scan window). Tip: clear a company's Last checked cell to "
     "give it a fresh first scan. 0 = always use the scan window.", False, None),
    ("watch_warn", "Watch: warn when a priority hasn't been scanned in (days)", 10,
     "Check Jobs warns you before a priority group's gap gets longer than the scan window.", False, None),
    ("watch_quiet", "Watch: flag a company as quiet after (days) with no matches", 30,
     "Its Last result says so; consider setting Active = No.", False, None),
    (H, "EMAIL CHECK", None, None, False, None),
    ("lookback_days", "First run: look back this many days", 60, "Also catches applications sent before you used the tracker.", False, None),
    ("followup_days", "Suggest follow-up after (days, no reply)", 10, "Sets Next Action to 'Follow up?'", False, None),
    ("ghost_days", "Mark Ghosted after (days, no reply)", 30, "0 = never", False, None),
    ("auto_add", "Add jobs found in confirmation emails", "Yes", "Yes/No", False, None),
    ("phone_keyword", "Phone capture: subject keyword", "job", "Email a link to yourself with this word in the subject. "
     "Add 'applied' to the subject if you already applied.", False, None),
    (H, "AI HELPER (free, runs on your Mac)", None, None, False, None),
    ("use_ai", "Use local AI", "Yes", "Yes/No. Needs the free Ollama app. Everything still works without it.", False, None),
    ("ai_model", "AI model", "qwen3:4b", "qwen3:4b is the default. If the AI is slow or times out, use qwen3:1.7b (faster, a bit "
     "less accurate). With 16GB+ RAM and Apple silicon, qwen3:8b is more accurate. Run Setup after changing.", False, None),
    ("copy_desc", "Copy posting to clipboard after adding", "Yes", "Yes/No. For pasting into an AI chat to tailor your resume.", False, None),
]
SETTING_LABELS = {s[1]: s[0] for s in SETTINGS if s[0] != H}
LEGACY_LABELS = {"Dealbreaker words": "dealbreakers",
                 "Watch: first check of a company, show postings from the last (days)": "watch_days",
                 "Watch: only show postings from the last (days)": "watch_days"}  # labels from earlier versions


def num(v, default=0.0):
    try:
        if v is None or v == "" or isinstance(v, bool):
            return default
        return float(str(v).replace(",", "").replace("$", "").replace("%", ""))
    except ValueError:
        return default


def yes(v) -> bool:
    return str(v or "").strip().lower() in ("yes", "y", "true", "1")


def split_list(v) -> list:
    return [x.strip() for x in str(v or "").split(",") if x.strip()]


def read_settings(wb) -> dict:
    vals = {}
    if "Settings" in wb.sheetnames:
        for row in wb["Settings"].iter_rows(min_row=1, max_col=2, values_only=True):
            lab = str(row[0] or "").strip() if row else ""
            key = SETTING_LABELS.get(lab) or LEGACY_LABELS.get(lab)
            if key and (key not in vals or vals[key] in (None, "")):
                vals[key] = row[1]
    out = {}
    for key, _, default, _, _, _ in SETTINGS:
        if key == H:
            continue
        v = vals.get(key)
        if v is None or (isinstance(v, str) and not v.strip()):
            v = default
        out[key] = v
    return out


AREA_DEFAULTS = []  # name, center ("City, ST" or a 2-letter state = whole state), radius miles, extra points. A new
#                    tracker starts with none: you add the metros you'd move to on the Areas tab.


def read_areas(wb) -> list:
    if "Areas" not in wb.sheetnames:
        return [{"name": n, "center": c, "radius": r, "bonus": b} for n, c, r, b in AREA_DEFAULTS]
    out = []
    for row in wb["Areas"].iter_rows(min_row=2, max_col=4, values_only=True):
        if row and row[1]:
            out.append({"name": str(row[0] or row[1]).strip(), "center": str(row[1]).strip(),
                        "radius": num(row[2], 35) if row[2] not in (None, "") else None, "bonus": num(row[3], 0)})
    return out


# ---------------------------------------------------------------- state
def load_state() -> dict:
    if STATE_PATH.exists():
        try:
            return json.loads(STATE_PATH.read_text())
        except Exception:
            pass
    return {}


def save_state(state: dict):
    STATE_PATH.write_text(json.dumps(state, indent=1))


# ---------------------------------------------------------------- Excel lock
def excel_lock_path() -> Path:
    return WB_PATH.parent / ("~$" + WB_PATH.name)


def wait_for_excel_closed():
    while excel_lock_path().exists():
        print("\n  The tracker is open in Excel. Save and close it (Cmd+S, then Cmd+W),")
        ans = input("  then press Enter. (Type 'force' if Excel is already closed.) ")
        if ans.strip().lower() == "force":
            try:
                excel_lock_path().unlink()
            except OSError:
                pass
            break


def today() -> dt.date:
    return dt.date.today()


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return v


# ---------------------------------------------------------------- Tracker
class Tracker:
    """Loads every job from the three job tabs into self.jobs (a list of dicts keyed by column key).
    Scripts change the dicts; save() re-sorts jobs into the right tab and rewrites the tabs."""

    def __init__(self):
        if not WB_PATH.exists():
            sys.exit(f"Can't find {WB_PATH.name} next to the app folder.")
        wait_for_excel_closed()
        import sheets
        self.wb = openpyxl.load_workbook(WB_PATH)
        sheets.ensure_structure(self.wb)
        self.s = read_settings(self.wb)
        self.areas = read_areas(self.wb)
        import geo
        self.geo_cfg = geo.config(self.s, self.areas)
        self.geo_sig = hashlib.md5(json.dumps(self.geo_cfg, sort_keys=True, default=str).encode()).hexdigest()[:6]
        self.jobs: list[dict] = []
        self.widths: dict = {}
        self.group_hidden = True
        for name in JOB_TABS + ["Jobs"]:  # "Jobs" = sheet name in the first version
            if name in self.wb.sheetnames:
                ws = self.wb[name]
                self.jobs += read_tab(ws)
                if name in JOB_TABS:
                    self._remember_layout(ws)

    def _remember_layout(self, ws):
        hdr = [c.value for c in ws[1]]
        for i, h in enumerate(hdr, 1):
            k = HEADER_TO_KEY.get(str(h or "").strip().lower())
            cd = ws.column_dimensions.get(get_column_letter(i))
            if k and cd is not None and cd.width and k not in self.widths and not cd.hidden:
                self.widths[k] = cd.width
            if k == "fit_score" and cd is not None and ws.title == "To Apply":
                self.group_hidden = bool(cd.hidden)

    def locate(self, job: dict):
        """Fill Area / Area Type / Miles for a job (only when its Location or your areas changed)."""
        import geo
        loc = str(job.get("location") or "").strip()
        key = f"{loc} #{self.geo_sig}"
        if job.get("geo_key") == key:
            return
        res = geo.classify(loc, self.geo_cfg) if loc else {"area": None, "area_type": None, "miles": None, "area_bonus": None}
        job.update(res)
        job["geo_key"] = key

    def next_id(self) -> int:
        return int(max([num(j.get("id"), 0) for j in self.jobs] + [0])) + 1

    def find_by_link(self, url):
        nu = norm_url(url)
        for j in self.jobs:
            if j.get("link") and norm_url(str(j["link"])) == nu:
                return j
        return None

    def add(self, job: dict) -> dict:
        job.setdefault("id", self.next_id())
        job.setdefault("date_added", today())
        self.jobs.append(job)
        return job

    def save(self):
        import scoring
        import sheets
        for j in self.jobs:
            normalize(j, self)
            self.locate(j)
        tabs = {t: [] for t in JOB_TABS}
        for j in self.jobs:
            tabs[tab_for(j)].append(j)
        s = self.s
        tabs["To Apply"].sort(key=lambda j: scoring.to_apply_sort_key(j, s))
        tabs["Applied"].sort(key=lambda j: (0 if j.get("next_action") else 1, -status_rank(j.get("status")),
                                            -_ord(j.get("last_update") or j.get("date_applied"))))
        tabs["Closed"].sort(key=lambda j: -_ord(j.get("last_update") or j.get("date_added")))
        if "Jobs" in self.wb.sheetnames:
            self.wb.remove(self.wb["Jobs"])
        for i, t in enumerate(JOB_TABS):
            sheets.write_job_tab(self.wb, t, tabs[t], i, self.widths, self.group_hidden)
        sheets.finalize(self.wb)
        save_workbook(self.wb)
        return {t: len(v) for t, v in tabs.items()}


def _ord(d):
    d = to_date(d)
    return d.toordinal() if isinstance(d, dt.date) else 0


def normalize(j: dict, tr: Tracker):
    if j.get("id") in (None, ""):
        j["id"] = tr.next_id()
    if not j.get("date_added"):
        j["date_added"] = today()
    st = str(j.get("status") or "").strip().title()
    st = st if st in STATUSES else ("" if not st else st)
    ap = "Yes" if yes(j.get("applied")) else "No"
    if st and st not in ("Saved", "Skipped") and st in STATUSES:
        ap = "Yes"
    if not st:
        st = "Applied" if ap == "Yes" else "Saved"
    if ap == "Yes" and st == "Saved":
        st = "Applied"
    if ap == "Yes" and not j.get("date_applied"):
        j["date_applied"] = today()
        if not j.get("resume") and not j.get("next_action"):
            j["next_action"] = "Add resume version"
    j["status"], j["applied"] = st, ap


def tab_for(j) -> str:
    if j.get("status") in CLOSED:
        return "Closed"
    if j.get("applied") == "Yes":
        return "Applied"
    return "To Apply"


def read_tab(ws) -> list:
    hdr = [HEADER_TO_KEY.get(str(c.value or "").strip().lower()) for c in ws[1]]
    jobs = []
    for row in ws.iter_rows(min_row=2):
        j = {}
        for k, c in zip(hdr, row):
            if not k or k in CALC:
                continue
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                continue
            if isinstance(v, str):
                v = v.strip() or None
            if k in DATE_KEYS:
                v = to_date(v)
            if k == "link" and not v and c.hyperlink is not None:
                v = c.hyperlink.target
            if v is not None:
                j[k] = v
        if j.get("company") or j.get("title") or j.get("link"):
            jobs.append(j)
    return jobs


def save_workbook(wb):
    wait_for_excel_closed()
    BACKUP_DIR.mkdir(exist_ok=True)
    if WB_PATH.exists():
        stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
        shutil.copy2(WB_PATH, BACKUP_DIR / f"Job_Tracker_{stamp}.xlsx")
        for f in sorted(BACKUP_DIR.glob("Job_Tracker_*.xlsx"))[:-BACKUPS_KEEP]:
            f.unlink()
    wb.calculation.fullCalcOnLoad = True
    # Write a complete new copy first, then swap it in, so a crash can't leave a half-written tracker.
    # The temporary name must NOT start with "." : macOS can keep such a file hidden in Finder after the rename.
    tmp = WB_PATH.with_name(WB_PATH.stem + " (saving).xlsx")
    wb.save(tmp)
    os.replace(tmp, WB_PATH)
    unhide(WB_PATH)
    old_tmp = WB_PATH.with_name(".saving_" + WB_PATH.name)  # left over from earlier versions
    if old_tmp.exists():
        try:
            old_tmp.unlink()
        except OSError:
            pass


def unhide(path):
    """Clear macOS's 'hidden' flag (no-op elsewhere)."""
    import stat
    try:
        st = os.stat(path)
        flags = getattr(st, "st_flags", 0)
        if flags & getattr(stat, "UF_HIDDEN", 0x8000):
            os.chflags(path, flags & ~stat.UF_HIDDEN)
    except (AttributeError, OSError):
        pass


def add_log(wb, sheet, values: list):
    ws = wb[sheet]
    r = ws.max_row + 1
    for i, v in enumerate(values, 1):
        c = ws.cell(r, i, v)
        if isinstance(v, (dt.date, dt.datetime)):
            c.number_format = "m/d/yy"
        if isinstance(v, str) and v.startswith("https://mail.google.com"):
            c.hyperlink = v
            c.value = "Open in Gmail"
            c.style = "Hyperlink"
    return r


def norm_url(u: str) -> str:
    u = (u or "").strip()
    u = re.sub(r"[?&](utm_[^=&]+|src|source|ref|refId|trackingId|trk|lipi|gh_src)=[^&]*", "", u)
    return u.rstrip("/?&").lower()


# ---------------------------------------------------------------- clipboard
def clipboard_get() -> str:
    try:
        return subprocess.run(["pbpaste"], capture_output=True, text=True, timeout=5).stdout
    except Exception:
        return ""


def clipboard_set(text: str):
    try:
        subprocess.run(["pbcopy"], input=text, text=True, timeout=5)
    except Exception:
        pass


# ---------------------------------------------------------------- local AI (Ollama)
OLLAMA = "http://localhost:11434"


OLLAMA_PATHS = ["/Applications/Ollama.app/Contents/Resources/ollama", "/opt/homebrew/bin/ollama", "/usr/local/bin/ollama"]
_server = None  # the background Ollama process this script started, if any


def _ollama_up() -> bool:
    try:
        return requests.get(f"{OLLAMA}/api/tags", timeout=2).ok
    except Exception:
        return False


def start_ollama() -> bool:
    """Start Ollama's background service (no app window). It's stopped again when the script exits."""
    global _server
    if _ollama_up():
        return True
    exe = shutil.which("ollama") or next((p for p in OLLAMA_PATHS if os.path.exists(p)), None)
    if not exe:
        return False
    import atexit
    import time
    try:
        _server = subprocess.Popen([exe, "serve"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                                   start_new_session=True)
    except OSError:
        return False
    atexit.register(stop_ollama)
    for _ in range(40):
        if _ollama_up():
            return True
        time.sleep(0.25)
    return False


def stop_ollama():
    global _server
    if _server is not None and _server.poll() is None:
        _server.terminate()
        try:
            _server.wait(timeout=5)
        except Exception:
            _server.kill()
    _server = None


class AI:
    """Thin wrapper around a local Ollama model. All calls fail soft (return None)."""

    def __init__(self, settings: dict):
        self.model = str(settings.get("ai_model") or "qwen3:4b").strip()
        self.enabled = False
        self.reason = "turned off in Settings"
        if not yes(settings.get("use_ai")):
            return
        if not start_ollama():
            self.reason = "Ollama isn't installed or wouldn't start (run Setup again)"
            return
        try:
            tags = requests.get(f"{OLLAMA}/api/tags", timeout=2).json()
        except Exception:
            self.reason = "Ollama didn't respond"
            return
        names = [m.get("name", "") for m in tags.get("models", [])]
        if not any(n == self.model or n.split(":")[0] == self.model or n.startswith(self.model) for n in names):
            self.reason = f"model {self.model} not downloaded (run Setup again)"
            return
        self.enabled = True
        self.reason = ""

    FIRST_TIMEOUT = 150  # the first request also loads the model into memory
    TIMEOUT = 75

    def json(self, prompt: str, schema: dict, system: str = "") -> dict | None:
        """One AI request. Never hangs: if the Mac is too slow, the AI is switched off for the rest of this run."""
        import time
        if not self.enabled:
            return None
        body = {
            "model": self.model,
            "messages": ([{"role": "system", "content": system}] if system else [])
            + [{"role": "user", "content": prompt}],
            "format": schema,
            "stream": False,
            "think": False,
            "keep_alive": "10m",
            "options": {"temperature": 0, "num_ctx": 6144},
        }
        limit = self.FIRST_TIMEOUT if not getattr(self, "calls", 0) else self.TIMEOUT
        self.calls = getattr(self, "calls", 0) + 1
        t0 = time.time()
        r = None
        for attempt in range(2):
            try:
                r = requests.post(f"{OLLAMA}/api/chat", json=body, timeout=limit)
                if r.status_code == 400 and "think" in r.text.lower():
                    body.pop("think", None)
                    continue
                r.raise_for_status()
                content = r.json()["message"]["content"]
                content = re.sub(r"<think>.*?</think>", "", content, flags=re.S).strip()
                self.last_secs = time.time() - t0
                return json.loads(content)
            except requests.exceptions.Timeout:
                self.enabled = False
                self.reason = f"took over {limit}s, too slow on this Mac right now"
                print(f"\n  (The AI took over {limit} seconds, so it's switched off for the rest of this run.\n"
                      f"   Everything still works with the rules. If this keeps happening, try the smaller model:\n"
                      f"   Settings tab > AI model > qwen3:1.7b, then run Setup.)")
                return None
            except Exception as e:
                if attempt == 1:
                    msg = ""
                    try:
                        msg = r.json().get("error", "")  # Ollama explains failures, e.g. not enough memory
                    except Exception:
                        msg = str(e)
                    self.last_error = (msg or type(e).__name__)[:200]
                    if not getattr(self, "_warned", False):
                        self._warned = True
                        print(f"\n  (AI request failed: {self.last_error}. Continuing with the rules.)")
                    return None
        return None


# ---------------------------------------------------------------- pay parsing
_NUM = r"(\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)"
PAY_RE = re.compile(
    r"(?:USD\s*)?\$\s?" + _NUM + r"\s*([kK])?(?:\s*(?:/|per|an|a)\s*(?:hour|hr|year|yr)\b)?"
    r"(?:\s*(?:-|–|—|to|and)\s*(?:USD\s*)?\$?\s?" + _NUM + r"\s*([kK])?)?"
    r"(?:\s*(?:USD|usd))?"
    r"(?:\s*(?:/|per|an|a)\s*(hour|hr|year|yr|annum|month|mo|week|wk)\b|\s*(hourly|annually|yearly))?",
    re.I,
)
# words that mean a dollar amount is NOT base pay
NOT_BASE = re.compile(r"sign[- ]?on|signing|bonus|stipend|allowance|reimburse|relocation|referral|401|match|"
                      r"per diem|tuition|equity|stock|commission", re.I)


def _to_num(s, k):
    v = float(s.replace(",", ""))
    return v * 1000 if k else v


def parse_pay(text: str):
    """Return (min, max, unit) where unit is 'Year' or 'Hour', or (None, None, None)."""
    if not text:
        return None, None, None
    best = None
    for m in PAY_RE.finditer(text):
        before = re.split(r"[.;:]\s|\n", text[max(0, m.start() - 45):m.start()])[-1]  # this sentence only
        after = text[m.end():m.end() + 25]
        extra_after = re.match(r"\s*(\+|plus\b|and\b|,|\(?\s*(?:not including|excluding|exclusive of)\b|(?:annual |target )?bonus[- ]eligible)", after, re.I)
        if NOT_BASE.search(before) or (not extra_after and NOT_BASE.search(after.split(".")[0])):
            continue
        lo = _to_num(m.group(1), m.group(2))
        hi = _to_num(m.group(3), m.group(4)) if m.group(3) else None
        if hi is not None and m.group(2) and not m.group(4) and hi < 1000:
            hi *= 1000  # "$80k-95" style
        if hi is not None and m.group(4) and not m.group(2) and lo < 1000:
            lo *= 1000
        unit_word = (m.group(5) or m.group(6) or "").lower()
        if unit_word.startswith(("hour", "hr")):
            unit = "Hour"
        elif unit_word.startswith(("mo",)):
            unit = "Month"
        elif unit_word.startswith(("week", "wk")):
            unit = "Week"
        elif unit_word:
            unit = "Year"
        else:
            unit = "Hour" if lo < 300 else "Year"
        if unit == "Year" and lo < 15000:
            continue
        if unit == "Hour" and not (7 <= lo <= 500):
            continue
        if hi is not None and hi < lo:
            hi = None
        if unit == "Month":
            lo, hi, unit = lo * 12, (hi * 12 if hi else None), "Year"
        if unit == "Week":
            lo, hi, unit = lo * 52, (hi * 52 if hi else None), "Year"
        rd = (lambda v: round(v, 2)) if unit == "Hour" else round
        cand = (rd(lo), rd(hi) if hi else rd(lo), unit)
        if best is None or (hi is not None and best[0] == best[1]):
            best = cand
            if hi is not None:
                break
    return best if best else (None, None, None)


def fmt_money(v) -> str:
    return f"${num(v):,.0f}"


def fmt_pay(pmin, pmax, unit) -> str:
    if not pmin and not pmax:
        return "not listed"
    u = "/hr" if str(unit).lower().startswith("h") or num(pmin or pmax) < 300 else "/yr"
    f = (lambda v: f"${num(v):,.2f}".replace(".00", "")) if u == "/hr" else (lambda v: f"${num(v):,.0f}")
    if pmin and pmax and num(pmin) != num(pmax):
        return f"{f(pmin)} - {f(pmax)}{u}"
    return f"{f(pmin or pmax)}{u}"


# ---------------------------------------------------------------- work mode
def detect_mode(location: str, title: str, desc: str) -> str:
    head = f"{title} {location}".lower()
    body = (desc or "")[:4000].lower()
    if "hybrid" in head:
        return "Hybrid"
    if re.search(r"\bremote\b", head) and "not remote" not in head:
        return "Remote"
    if re.search(r"\bhybrid\b", body):
        return "Hybrid"
    if re.search(r"fully remote|100% remote|remote-first|work from anywhere|this (role|position) is remote|"
                 r"\bremote (role|position|opportunity|job)\b", body):
        return "Remote"
    if re.search(r"\b(on-?site|onsite|in[- ]office|in the office|in[- ]person)\b", body):
        return "On-site"
    return ""


def safe_name(s: str, n=40) -> str:
    return re.sub(r"[^\w\- ]+", "", str(s or "")).strip().replace(" ", "_")[:n] or "job"
